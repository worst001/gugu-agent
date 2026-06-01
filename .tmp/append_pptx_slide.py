from __future__ import annotations

import sys
from copy import deepcopy
from pathlib import Path

from pptx import Presentation


DRAWING_TAGS = {
    "sp",
    "grpSp",
    "graphicFrame",
    "cxnSp",
    "pic",
    "contentPart",
}


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def clear_slide_shapes(slide) -> None:
    sp_tree = slide.shapes._spTree
    for element in list(sp_tree):
        if local_name(element.tag) in DRAWING_TAGS:
            sp_tree.remove(element)


def append_slide_shapes(source_slide, dest_slide) -> None:
    sp_tree = dest_slide.shapes._spTree
    for shape in source_slide.shapes:
        sp_tree.insert_element_before(deepcopy(shape._element), "p:extLst")


def main() -> int:
    if len(sys.argv) != 4:
        print("Usage: append_pptx_slide.py <base.pptx> <one-slide.pptx> <output.pptx>")
        return 2

    base_path = Path(sys.argv[1])
    one_slide_path = Path(sys.argv[2])
    output_path = Path(sys.argv[3])

    base = Presentation(base_path)
    addition = Presentation(one_slide_path)
    if len(addition.slides) != 1:
        raise ValueError(f"Expected one slide in {one_slide_path}, found {len(addition.slides)}")

    source_slide = addition.slides[0]
    layout = base.slide_layouts[-1]
    dest_slide = base.slides.add_slide(layout)
    clear_slide_shapes(dest_slide)
    append_slide_shapes(source_slide, dest_slide)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    base.save(output_path)
    print(f"Saved {output_path}")
    print(f"Slides: {len(base.slides)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
